import time
import streamlit as st
import cv2
import keras
import numpy as np
from openpyxl import load_workbook
from collections import deque
from PIL import Image
from streamlit_TTS import auto_play, text_to_speech, text_to_audio
from gtts.lang import tts_langs

# Ensure the camera feed is linked to the model for predictions

# Define missing variables
max_frames = 30  # Adjust as needed
input_size = (128, 128)  # Adjust as needed
predict_every_n_frames = 60  # Adjust as needed

# Ensure model is loaded
if 'model' not in st.session_state:
    @st.cache_resource
    def load_model():
        try:
            return keras.models.load_model('/Users/fayadh/code/SaadAlshahrani/arabic-sign-language-translator/arabic_sign_language_translator/New/3_words_model_new.keras')
        except Exception as e:
            st.error(f"Failed to load model: {e}")
            return None
    st.session_state.model = load_model()

model = st.session_state.model

# Ensure label map is loaded
if 'signid_map' not in st.session_state:
    @st.cache_resource
    def load_label_map(excel_path):
        wb = load_workbook(excel_path)
        ws = wb.active
        sign_ids = [str(row[1]).zfill(4) for row in ws.iter_rows(min_row=2, values_only=True)]
        return {idx: sid for idx, sid in enumerate(sign_ids)}
    st.session_state.signid_map = load_label_map("KARSL-100_Labels.xlsx")

signid_map = st.session_state.signid_map

# Ensure placeholders are defined
if 'video_placeholder' not in st.session_state:
    st.session_state.video_placeholder = st.empty()
if 'prediction_placeholder' not in st.session_state:
    st.session_state.prediction_placeholder = st.empty()

video_placeholder = st.session_state.video_placeholder
prediction_placeholder = st.session_state.prediction_placeholder

if st.session_state.camera_running:
    cap = cv2.VideoCapture(0)
    frame_count = 0
    while st.session_state.camera_running:
        ret, frame = cap.read()
        if not ret:
            st.error("Failed to access camera. Please check your webcam connection.")
            break

        # Preprocess the frame
        frame_rgb = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
        resized_frame = cv2.resize(frame_rgb, input_size)
        normalized_frame = resized_frame / 255.0
        st.session_state.frame_buffer.append(normalized_frame)

        # Perform prediction at specified intervals
        if len(st.session_state.frame_buffer) == max_frames and frame_count % predict_every_n_frames == 0:
            input_batch = np.expand_dims(np.array(st.session_state.frame_buffer), axis=0)
            preds = model.predict(input_batch, verbose=0)
            pred_class = np.argmax(preds)
            st.session_state.prediction = signid_map[pred_class]
            st.session_state.confidence = np.max(preds)

        # Display the camera feed
        with video_placeholder.container():
            st.image(frame, channels="BGR", use_container_width=True)

        # Display prediction and confidence
        if st.session_state.prediction:
            with prediction_placeholder.container():
                st.markdown(f'<p class="prediction-text">Predicted Sign: <strong>{st.session_state.prediction}</strong></p>', unsafe_allow_html=True)
                st.markdown(f'<p class="confidence-text">Confidence: <strong>{st.session_state.confidence:.2%}</strong></p>', unsafe_allow_html=True)

        frame_count += 1
    cap.release()
else:
    # Display message when the camera is off
    with prediction_placeholder.container():
        st.markdown('<p class="Cam-OFF-Text">Camera is off. Start the camera to begin recognition.</p>', unsafe_allow_html=True)
